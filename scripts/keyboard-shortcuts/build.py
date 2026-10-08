import json
from datetime import date
import os
import re
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

WORK, DOTFILES, OUTPUT = sys.argv[1:4]
NVIM_CFG = os.path.join(WORK, "cfg/nvim")

ARIAL = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BORDER = Border(bottom=Side(style="thin", color="D1D5DB"))
GREY = Font(name=ARIAL, size=10, italic=True, color="9CA3AF")


def style_sheet(ws, headers, rows, widths, wrap_cols, grey_when=None):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22
    for i, row in enumerate(rows, start=1):
        ws.append([i] + row)
    last = ws.max_row
    call_col = headers.index("Your call") + 1
    for r in ws.iter_rows(min_row=2, max_row=last):
        for c in r:
            c.font = Font(name=ARIAL, size=10)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap_cols)
        r[1].font = Font(name=ARIAL, size=10, bold=True)
        r[call_col - 1].fill = INPUT_FILL
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    call = ws.cell(1, call_col).column_letter
    dv = DataValidation(type="list", formula1='"Keep,Remove,Rebind"', allow_blank=True)
    dv.prompt = "Keep, Remove or Rebind"
    ws.add_data_validation(dv)
    dv.add(f"{call}2:{call}{last}")
    end = ws.cell(1, len(headers)).column_letter
    rng = f"A2:{end}{last}"
    if grey_when:
        col, values = grey_when
        for v in values:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col}2="{v}"'], font=GREY))
    for word, color in (("Remove", "FDE2E2"), ("Rebind", "E0ECFF"), ("Keep", "E3F4E1")):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${call}2="{word}"'], fill=PatternFill("solid", fgColor=color)))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{end}{last}"
    return last, call


OMARCHY_GROUPS = OrderedDict([
    ("tiling", "Windows & workspaces"), ("applications", "Apps"), ("utilities", "Menus, capture & system"),
    ("media", "Volume, media & brightness"), ("clipboard", "Clipboard"), ("voxtype", "Dictation"),
])
HYPR_KEYS = {
    "mouse:272": "Left-drag", "mouse:273": "Right-drag", "mouse_down": "Scroll down", "mouse_up": "Scroll up",
    "code:20": "-", "code:21": "=", "RETURN": "Enter", "BACKSPACE": "Backspace", "SLASH": "/", "COMMA": ",",
    "PERIOD": ".", "SPACE": "Space", "ESCAPE": "Esc", "PRINT": "Print Screen",
}
for n in range(10, 20):
    HYPR_KEYS[f"code:{n}"] = str((n - 9) % 10)


def hypr_norm(keys):
    return re.sub(r"\s+", "", keys).upper()


def hypr_pretty(keys):
    out = []
    for p in [p.strip() for p in keys.split("+")]:
        if p in HYPR_KEYS or p.upper() in HYPR_KEYS:
            out.append(HYPR_KEYS.get(p, HYPR_KEYS.get(p.upper())))
        elif p.upper() in ("SUPER", "SHIFT", "CTRL", "ALT"):
            out.append(p.capitalize())
        elif p.startswith("XF86"):
            out.append(p[4:] + " key")
        elif len(p) == 1:
            out.append(p.upper())
        else:
            out.append(p.capitalize() if p.isupper() else p)
    return " + ".join(out)


def omarchy_rows():
    records = [json.loads(l) for l in open(os.path.join(WORK, "hypr-binds.jsonl"))]
    unbound = {hypr_norm(r["keys"]) for r in records if r["op"] == "unbind"}
    yours = OrderedDict()
    for r in records:
        if r["op"] == "bind" and r["file"] == "yours":
            yours.setdefault(hypr_norm(r["keys"]), r)
    rows = OrderedDict()
    for r in records:
        if r["op"] != "bind" or r["file"] == "yours":
            continue
        k = (r["file"], hypr_norm(r["keys"]))
        if k in rows:
            if r["desc"] and r["desc"] not in rows[k]["desc"]:
                rows[k]["desc"] += "; " + r["desc"]
            continue
        rows[k] = dict(r)
    table = []
    for (file, norm), r in rows.items():
        if norm in yours:
            status, note = "Replaced by yours", "Now: " + (yours[norm]["desc"] or "custom")
        elif norm in unbound:
            status, note = "Removed by you", ""
        else:
            status, note = "Active", ""
        table.append([hypr_pretty(r["keys"]), r["desc"], OMARCHY_GROUPS.get(file, file), "Omarchy", status, None, note])
    defaults = {n for (_, n) in rows}
    for norm, r in yours.items():
        table.append([hypr_pretty(r["keys"]), r["desc"], "Yours", "You", "Active", None,
                      "Replaces an Omarchy default" if norm in defaults else ""])
    order = list(OMARCHY_GROUPS.values()) + ["Yours"]
    table.sort(key=lambda row: order.index(row[2]))
    return table


MODES = OrderedDict([("n", "Normal"), ("x", "Visual"), ("v", "Visual"), ("s", "Select"), ("o", "Motion"),
                     ("i", "Insert"), ("c", "Command line"), ("t", "Terminal")])
VIM_SPECIAL = {"Space": "Space", "CR": "Enter", "BS": "Backspace", "Esc": "Esc", "Tab": "Tab", "NL": "Ctrl+J",
               "lt": "<", "Bslash": "\\", "Bar": "|", "Up": "Up", "Down": "Down", "Left": "Left", "Right": "Right",
               "Nop": "nothing"}


def vim_tokens(lhs):
    return re.findall(r"<[^<>]+>|.", lhs)


def vim_pretty(lhs):
    out = []
    for t in vim_tokens(lhs):
        if t.startswith("<") and t.endswith(">") and len(t) > 2:
            inner = t[1:-1]
            if inner in VIM_SPECIAL:
                out.append(VIM_SPECIAL[inner])
                continue
            m = re.match(r"^([CMSAD])-(.+)$", inner, re.I)
            if m:
                mod = {"C": "Ctrl", "M": "Alt", "A": "Alt", "S": "Shift", "D": "Cmd"}[m.group(1).upper()]
                key = VIM_SPECIAL.get(m.group(2), m.group(2))
                out.append(f"{mod}+{key.upper() if len(key) == 1 else key}")
                continue
            out.append(inner)
        else:
            out.append(t)
    return " ".join(out)


def vim_norm(lhs):
    s = lhs.replace("<leader>", "<Space>").replace("<Leader>", "<Space>")
    s = re.sub(r"<[^<>]+>", lambda m: m.group(0).upper(), s)
    return s.replace("<C-J>", "<NL>")


def mode_text(modes):
    seen = []
    for m in MODES:
        if m in modes and MODES[m] not in seen:
            seen.append(MODES[m])
    return ", ".join(seen)


def readable_desc(x):
    desc = x["desc"] or ""
    rhs = x["rhs"] or ""
    if desc == "which_key_ignore":
        return "Alias of Ctrl+/ (terminals send Ctrl+/ as Ctrl+_)"
    if desc.startswith(":help "):
        return "Neovim default (see " + desc + ")"
    if desc:
        return desc
    known = {"<lt>gv": "Indent left, keep selection", ">gv": "Indent right, keep selection"}
    if rhs in known:
        return known[rhs]
    if rhs.endswith("<C-G>u"):
        return "Undo breakpoint after punctuation"
    return rhs or "(no description)"


def builtin_source(x):
    d = x.get("defined_in") or ""
    if "_core/defaults" in d or (x["desc"] or "").startswith(":help") or x["lhs"] == "<C-W><C-D>" or x["sid"] == -8:
        return "Neovim built-in"
    if x["script"].endswith("matchit.vim"):
        return "Neovim built-in (matchit)"
    return None


LAZYVIM_GROUPS = {"b": "Buffers", "c": "Code", "d": "Debug", "f": "Files / find", "g": "Git", "q": "Quit / session",
                  "s": "Search", "u": "UI toggles", "w": "Windows", "x": "Diagnostics / quickfix"}
CUSTOM_GROUPS = {"c": "Code", "d": "Document / debug", "f": "Find", "g": "Git", "h": "Git hunks", "q": "Session",
                 "r": "Rename / restart", "s": "Search", "t": "Toggles", "w": "Workspace"}


def vim_category(lhs, source, groups):
    if source.startswith("mini.pairs") or source == "nvim-autopairs":
        return "Auto-pairs while typing"
    if source.startswith("Neovim built-in"):
        return "Neovim built-in"
    if lhs.startswith("<Space>"):
        rest = lhs[len("<Space>"):]
        if rest.startswith("<Tab>"):
            return "Space: Tabs"
        if rest[:1] in groups and len(rest) > 1:
            return "Space: " + groups[rest[:1]]
        return "Space: Other"
    if lhs[:1] in "[]":
        return "Jump next / previous"
    if lhs.startswith("g"):
        return "g commands"
    if lhs.startswith(("<C-", "<M-")):
        return "Ctrl / Alt keys"
    return "Other"


def lazyvim_rows():
    data = json.load(open(os.path.join(WORK, "lazyvim-maps.json")))
    merged = OrderedDict()
    for x in data["maps"]:
        if x["lhs"].startswith(("<Plug>", "<SNR>")):
            continue
        src = builtin_source(x)
        if not src:
            d = x.get("defined_in") or ""
            m = re.search(r"/lazy/([^/]+)/", d)
            if x["plugin"]:
                src = x["plugin"]
            elif "/.config/nvim/" in d:
                src = "Yours"
            elif m:
                src = m.group(1)
            elif "MiniPairs" in (x["rhs"] or ""):
                src = "mini.pairs (auto-pairs)"
            else:
                src = "LazyVim"
        merged.setdefault((x["lhs"], readable_desc(x), src), set()).add(x["mode"])
    lsp = OrderedDict()
    for k in data["lsp"]:
        lhs = k["lhs"].replace("<leader>", "<Space>").replace("<c-k>", "<C-K>").replace("<a-n>", "<M-n>").replace("<a-p>", "<M-p>")
        modes = k["mode"] if isinstance(k["mode"], list) else [k["mode"]]
        lsp[(lhs, tuple(modes))] = k["desc"]
    for (lhs, modes), desc in lsp.items():
        merged.setdefault((lhs, desc, "LazyVim (LSP)"), set()).update(modes)
    rows = []
    for (lhs, desc, src), modes in merged.items():
        cat = "Code (LSP, in code files)" if src == "LazyVim (LSP)" else vim_category(lhs, src, LAZYVIM_GROUPS)
        rows.append([vim_pretty(lhs), mode_text(modes), desc, cat, src, None, None])
    order = ["Yours", "LazyVim", "LazyVim (LSP)"]

    def key(r):
        s = r[4]
        rank = order.index(s) if s in order else 5 if s.startswith("Neovim built-in") else 6 if s.startswith("mini.pairs") else 4
        return (rank, r[3], r[0].lower())
    rows.sort(key=key)
    return rows


def index_custom_config():
    index = {}
    pattern_set = re.compile(r"""(?:vim\.keymap\.set|\bmap)\(\s*(\{[^}]*\}|'[^']*'|"[^"]*")\s*,\s*(?:'([^']*)'|"([^"]*)")""")
    files = []
    for root, _, names in os.walk(NVIM_CFG):
        for f in names:
            if f.endswith(".lua"):
                files.append(os.path.join(root, f))
    files.sort(key=lambda p: (0 if "/lua/config/" in p else 1, p))
    for path in files:
        rel = os.path.relpath(path, NVIM_CFG)
        for m in pattern_set.finditer(open(path).read()):
            index.setdefault(vim_norm(m.group(2) or m.group(3)), rel)
    return index


def plugin_spec_file(plugin):
    base = plugin.replace(".nvim", "").replace("nvim-", "")
    for root, _, names in os.walk(os.path.join(NVIM_CFG, "lua/custom/plugins")):
        for f in sorted(names):
            text = open(os.path.join(root, f)).read()
            if plugin in text or ("/" + base) in text:
                return os.path.relpath(os.path.join(root, f), NVIM_CFG)
    return None


def buffer_local_custom():
    rows = []
    lsp = open(os.path.join(NVIM_CFG, "lua/custom/plugins/lsp.lua")).read()
    for m in re.finditer(r"""nmap\(\s*'([^']+)'.*?,\s*'([^']*)'\s*\)""", lsp, re.S):
        rows.append((m.group(1), "n", m.group(2), "lua/custom/plugins/lsp.lua", "Code (LSP, in code files)"))
    git = open(os.path.join(NVIM_CFG, "lua/custom/plugins/gitsigns.lua")).read()
    for m in re.finditer(r"""map\(\s*(\{[^}]*\}|'[^']*')\s*,\s*'([^']+)'(.*?)desc\s*=\s*'([^']*)'""", git, re.S):
        modes = re.findall(r"'(\w)'", m.group(1))
        rows.append((m.group(2), "".join(modes), m.group(4), "lua/custom/plugins/gitsigns.lua", "Git hunks (in git repos)"))
    return rows


def custom_nvim_rows():
    data = json.load(open(os.path.join(WORK, "nvim-maps.json")))
    index = index_custom_config()
    motions = open(os.path.join(NVIM_CFG, "lua/config/vim-motions.lua")).read()
    motion_lhs = {vim_norm(m) for m in re.findall(r"""vim\.keymap\.set\('n',\s*(?:'([^']*)'|"([^"]*)")""", motions) for m in [m[0] or m[1]]}
    merged = OrderedDict()
    for x in data["maps"]:
        if x["lhs"].startswith(("<Plug>", "<SNR>")):
            continue
        d = x.get("defined_in") or ""
        norm = vim_norm(x["lhs"])
        src = builtin_source(x)
        if not src and "autopairs" in (x["desc"] or "").lower():
            src = "nvim-autopairs (plugin default)"
        if not src and d.startswith("/usr/share/nvim/runtime") and norm in index:
            src = "Yours: " + index[norm]
        if not src and "/cfg/nvim/" in d:
            src = "Yours: " + d.split("/cfg/nvim/")[1]
        if not src and x["plugin"]:
            spec = plugin_spec_file(x["plugin"])
            src = ("Yours: " + spec) if spec else x["plugin"] + " (plugin default)"
        if not src and not d and norm in index and not x["script"].endswith(".vim"):
            src = "Yours: " + index[norm]
        if not src:
            m = re.search(r"/lazy/([^/]+)/", d)
            if m:
                src = m.group(1) + " (plugin default)"
            elif x["script"].endswith(("plugs.vim", "maps.vim")):
                src = "vim-visual-multi (plugin default)"
            elif x["script"].endswith("nvim-surround.lua"):
                src = "nvim-surround (plugin default)"
            elif x["script"].endswith("plenary.vim"):
                src = "plenary.nvim (plugin default)"
            else:
                src = "Plugin default"
        merged.setdefault((x["lhs"], readable_desc(x), src), set()).add(x["mode"])
    for lhs, modes, desc, rel, _ in buffer_local_custom():
        lhs = lhs.replace("<leader>", "<Space>")
        merged.setdefault((lhs, desc, "Yours: " + rel), set()).update(modes)

    rows = []
    for (lhs, desc, src), modes in merged.items():
        note = None
        if src == "Yours: lua/config/vim-motions.lua" or (src.startswith("Yours") and vim_norm(lhs) in motion_lhs and src.endswith("vim-motions.lua")):
            cat = "Vim motion (description only)"
            note = "Maps the key to itself so it shows up in keymap search. Removing it only drops the description."
        elif src.endswith("lsp.lua"):
            cat = "Code (LSP, in code files)"
        elif src.endswith("gitsigns.lua"):
            cat = "Git hunks (in git repos)"
        else:
            cat = vim_category(lhs, src, CUSTOM_GROUPS)
        rows.append([vim_pretty(lhs), mode_text(modes), desc, cat, src, None, note])

    def key(r):
        s = r[4]
        if s.startswith("Yours") and r[3] != "Vim motion (description only)":
            rank = 0
        elif s.startswith("Yours"):
            rank = 1
        elif s.startswith("Neovim built-in"):
            rank = 3
        else:
            rank = 2
        return (rank, s, r[3], r[0].lower())
    rows.sort(key=key)
    return rows


HERDR_DESC = {
    "prefix": "The prefix key every other herdr shortcut starts with", "help": "Show keybinding help",
    "settings": "Open settings", "detach": "Detach from the session (it keeps running)", "reload_config": "Reload config.toml",
    "open_notification_target": "Jump to the pane behind the latest notification", "workspace_picker": "Open the workspace picker",
    "goto": "Go-to picker", "new_workspace": "New workspace", "new_worktree": "New git worktree workspace",
    "open_worktree": "Open an existing worktree", "remove_worktree": "Remove a worktree (asks first)",
    "rename_workspace": "Rename workspace", "close_workspace": "Close workspace", "previous_workspace": "Previous workspace",
    "next_workspace": "Next workspace", "previous_agent": "Previous agent", "next_agent": "Next agent",
    "focus_agent": "Focus agent by number", "remote_image_paste": "Paste an image over herdr --remote", "new_tab": "New tab",
    "rename_tab": "Rename tab", "previous_tab": "Previous tab", "next_tab": "Next tab",
    "move_tab_previous": "Move tab toward the front", "move_tab_next": "Move tab toward the back",
    "switch_tab": "Switch to tab 1-9", "switch_workspace": "Switch to workspace by number", "close_tab": "Close tab",
    "rename_pane": "Rename pane", "edit_scrollback": "Open the pane's scrollback in your editor",
    "focus_pane_left": "Focus pane left", "focus_pane_down": "Focus pane below", "focus_pane_up": "Focus pane above",
    "focus_pane_right": "Focus pane right", "cycle_pane_next": "Next pane", "cycle_pane_previous": "Previous pane",
    "last_pane": "Last-used pane", "split_vertical": "Split side by side", "split_horizontal": "Split top / bottom",
    "close_pane": "Close pane", "zoom": "Zoom pane (fullscreen inside herdr)", "resize_mode": "Enter resize mode",
    "resize_pane_left": "Resize pane left", "resize_pane_down": "Resize pane down", "resize_pane_up": "Resize pane up",
    "resize_pane_right": "Resize pane right", "toggle_sidebar": "Toggle the sidebar",
    "navigate_workspace_up": "Navigate mode: workspace up", "navigate_workspace_down": "Navigate mode: workspace down",
    "navigate_pane_left": "Navigate mode: pane left", "navigate_pane_down": "Navigate mode: pane down",
    "navigate_pane_up": "Navigate mode: pane up", "navigate_pane_right": "Navigate mode: pane right",
}
H_ORDER = ["Session & app", "Workspaces", "Tabs", "Panes", "Agents", "Navigate mode", "Yours"]


def herdr_group(action):
    if action.startswith("navigate_"):
        return "Navigate mode"
    if "worktree" in action or "workspace" in action or action == "goto":
        return "Workspaces"
    if "tab" in action and "pane" not in action:
        return "Tabs"
    if "pane" in action or action in ("split_vertical", "split_horizontal", "zoom", "resize_mode"):
        return "Panes"
    if "agent" in action:
        return "Agents"
    return "Session & app"


def herdr_pretty(binding):
    if not binding:
        return ""
    parts = []
    for p in binding.split("+"):
        if p == "prefix":
            parts.append("Ctrl+B ▸")
            continue
        p = {"minus": "-", "tab": "Tab", "up": "Up", "down": "Down", "left": "Left", "right": "Right"}.get(p, p)
        parts.append(p if p == "1..9" else p.upper() if len(p) == 1 else p.capitalize())
    return "+".join(parts).replace("▸+", "▸ ")


def herdr_rows():
    rows = []
    for line in open(os.path.join(WORK, "herdr-keys.tsv")):
        action, binding, note = (line.rstrip("\n").split("\t") + ["", ""])[:3]
        rows.append([herdr_pretty(binding), action, HERDR_DESC.get(action, action.replace("_", " ").capitalize()),
                     herdr_group(action), "herdr", "Active" if binding else "Not bound", None, note.strip() or None])
    rows.append([herdr_pretty("prefix+alt+c"), "keys.command", "Run ~/.local/bin/herdr-tabs", "Yours", "You", "Active",
                 None, "From ~/.config/herdr/config.toml"])
    rows.sort(key=lambda r: (r[5] != "Active", H_ORDER.index(r[3])))
    return rows


SKHDRC = os.path.join(DOTFILES, "yabai/.skhdrc")
SKHD_KEYS = {"0x2F": ".", "0x2B": ",", "left": "Left", "right": "Right", "up": "Up", "down": "Down"}
DIRECTIONS = {"west": "left", "east": "right", "north": "up", "south": "down", "prev": "previous desktop", "next": "next desktop",
              "left": "previous desktop", "right": "next desktop"}


def skhd_pretty(chord):
    mods, _, key = chord.partition(" - ")
    parts = [m.strip().capitalize() for m in mods.split("+") if m.strip()]
    key = key.strip()
    parts.append(SKHD_KEYS.get(key, key.upper() if len(key) == 1 else key))
    return " + ".join(parts)


def skhd_describe(command):
    m = re.search(r"window --focus (\w+)", command)
    if m:
        return "Focus window " + DIRECTIONS[m.group(1)], "Focus"
    m = re.search(r"window --warp (\w+)", command)
    if m:
        return "Move window " + DIRECTIONS[m.group(1)], "Move windows"
    m = re.search(r"window --space (\w+)", command)
    if m:
        target = m.group(1)
        where = "desktop " + target if target.isdigit() else DIRECTIONS.get(target, target)
        return "Move window to " + where + " and follow", "Desktops"
    m = re.match(r"\S*space-follow (\w+)$", command.strip())
    if m:
        target = m.group(1)
        return ("Go to desktop " + target if target.isdigit() else "Go to " + DIRECTIONS.get(target, target)), "Desktops"
    if "zoom-fullscreen" in command:
        return "Toggle zoom-fullscreen", "Layout"
    if "--toggle float" in command:
        return "Float window and centre it", "Layout"
    m = re.search(r"space --layout (\w+)", command)
    if m:
        return {"bsp": "Tiling layout (bsp)", "stack": "Stacked layout"}.get(m.group(1), m.group(1)), "Layout"
    if "--balance" in command:
        return "Balance: make all windows equal size", "Layout"
    return command, "Other"


def skhd_keyboard(chord):
    mods = chord.partition(" - ")[0]
    if "ctrl" in mods and "alt" in mods and "cmd" in mods:
        return "Lily58"
    if "ctrl" in mods and "cmd" in mods:
        return "Omarchy-style"
    return "MacBook"


def yabai_rows():
    rows = []
    for line in open(SKHDRC):
        text = line.strip()
        active = not text.startswith("#")
        body = text.lstrip("#").strip()
        if not re.search(r" : .*(yabai|space-follow)", body) or not re.match(r"^[a-z0-9 +]+ - \S+\s*:", body):
            continue
        chord, _, command = body.partition(" : ")
        desc, group = skhd_describe(command)
        rows.append([skhd_pretty(chord), desc, group, skhd_keyboard(chord), "Active" if active else "Commented out", None, None])
    order = ["Focus", "Move windows", "Layout", "Desktops", "Other"]
    rows.sort(key=lambda r: (r[4] != "Active", order.index(r[2])))
    return rows


wb = Workbook()
summary = wb.active
summary.title = "Summary"

tabs = []

ws = wb.create_sheet("Omarchy")
last, call = style_sheet(ws, ["#", "Shortcut", "What it does", "Group", "Set by", "Status", "Your call", "Notes"],
                         omarchy_rows(), {"A": 5, "B": 30, "C": 46, "D": 26, "E": 10, "F": 18, "G": 12, "H": 34}, (3, 8),
                         grey_when=("F", ["Removed by you", "Replaced by yours"]))
tabs.append(("Omarchy", "D", last, call, "Omarchy / Hyprland, this machine. Grey rows: defaults you already removed or replaced."))

for title, rows, note in (
    ("LazyVim", lazyvim_rows(), "LazyVim on this machine (~/.config/nvim). Space is the leader key."),
    ("Nvim", custom_nvim_rows(), "Your own Neovim config (~/dotfiles/nvim, the Mac setup). Space is the leader key."),
):
    ws = wb.create_sheet(title)
    last, call = style_sheet(ws, ["#", "Keys", "Mode", "What it does", "Category", "Comes from", "Your call", "Notes"],
                             rows, {"A": 5, "B": 22, "C": 18, "D": 42, "E": 28, "F": 34, "G": 12, "H": 40}, (3, 4, 8))
    tabs.append((title, "E", last, call, note))

ws = wb.create_sheet("herdr")
last, call = style_sheet(ws, ["#", "Keys", "Setting name", "What it does", "Group", "Set by", "Status", "Your call", "Notes"],
                         herdr_rows(), {"A": 5, "B": 24, "C": 24, "D": 44, "E": 16, "F": 9, "G": 11, "H": 12, "I": 44}, (4, 9),
                         grey_when=("G", ["Not bound"]))
tabs.append(("herdr", "E", last, call, "herdr 0.9.1 defaults plus your config. Every shortcut starts with Ctrl+B (shown as \"Ctrl+B ▸\")."))

ws = wb.create_sheet("Yabai")
last, call = style_sheet(ws, ["#", "Shortcut", "What it does", "Group", "Keyboard", "Status", "Your call", "Notes"],
                         yabai_rows(), {"A": 5, "B": 32, "C": 40, "D": 16, "E": 16, "F": 16, "G": 12, "H": 40}, (3, 8),
                         grey_when=("F", ["Commented out"]))
tabs.append(("Yabai", "D", last, call, "skhd bindings for yabai on the Mac (~/dotfiles/yabai/.skhdrc). Grey rows are commented out."))

summary.column_dimensions["A"].width = 36
for col in "BCDE":
    summary.column_dimensions[col].width = 12
summary.append(["How to use this"])
summary["A1"].font = Font(name=ARIAL, bold=True, size=13)
notes = [
    "One tab per tool. Fill in the yellow \"Your call\" column: Keep, Remove or Rebind.",
    "Example: Omarchy | Super + Shift + Y | YouTube | Your call = Remove | Notes = never use it.",
] + [f"{t}: {n}" for t, _, _, _, n in tabs] + [
    f"Generated {date.today().isoformat()} by scripts/keyboard-shortcuts/refresh from the live Hyprland, LazyVim, Neovim, herdr and skhd configs.",
]
for n in notes:
    summary.append([n])
for r in range(2, 2 + len(notes)):
    summary.cell(r, 1).font = Font(name=ARIAL, size=10)

row = 2 + len(notes) + 1
for sheet, group_col, last, call, _ in tabs:
    summary.cell(row, 1, f"{sheet} — by {wb[sheet].cell(1, ord(group_col) - 64).value.lower()}")
    for c, h in zip(range(2, 6), ["Shortcuts", "Keep", "Remove", "Rebind"]):
        summary.cell(row, c, h)
    for c in range(1, 6):
        summary.cell(row, c).font = Font(name=ARIAL, bold=True, color="FFFFFF")
        summary.cell(row, c).fill = HEADER_FILL
    start = row + 1
    values = []
    for r in range(2, last + 1):
        v = wb[sheet].cell(r, ord(group_col) - 64).value
        if v not in values:
            values.append(v)
    rng = lambda col: f"'{sheet}'!${col}$2:${col}${last}"
    r = start
    for v in values:
        summary.cell(r, 1, v)
        summary.cell(r, 2, f'=COUNTIFS({rng(group_col)},$A{r})')
        summary.cell(r, 3, f'=COUNTIFS({rng(group_col)},$A{r},{rng(call)},"Keep")')
        summary.cell(r, 4, f'=COUNTIFS({rng(group_col)},$A{r},{rng(call)},"Remove")')
        summary.cell(r, 5, f'=COUNTIFS({rng(group_col)},$A{r},{rng(call)},"Rebind")')
        r += 1
    summary.cell(r, 1, "Total")
    for c, col in zip(range(2, 6), "BCDE"):
        summary.cell(r, c, f"=SUM({col}{start}:{col}{r - 1})")
    for rr in summary.iter_rows(min_row=start, max_row=r):
        for c in rr:
            c.font = Font(name=ARIAL, size=10, bold=(c.row == r))
            c.border = BORDER
    row = r + 2

wb.calculation.fullCalcOnLoad = True
wb.save(OUTPUT)
for t in tabs:
    print(t[0], t[2] - 1, "rows")
