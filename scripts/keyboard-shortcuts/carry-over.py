import sys

from openpyxl import load_workbook

old_path, new_path = sys.argv[1], sys.argv[2]

COLUMNS = {
    "Omarchy": ("Shortcut", "What it does"),
    "LazyVim": ("Keys", "What it does"),
    "Nvim": ("Keys", "What it does"),
    "herdr": ("Keys", "Setting name"),
    "Yabai": ("Shortcut", "What it does"),
}
GENERATED_NOTE_PREFIXES = ("Now: ", "Replaces an Omarchy default", "Maps the key to itself", "From ~/.config/herdr")


def header_index(ws):
    return {c.value: i for i, c in enumerate(ws[1])}


old = load_workbook(old_path)
new = load_workbook(new_path)

for sheet, (key_col, desc_col) in COLUMNS.items():
    if sheet not in old.sheetnames:
        continue
    o, n = old[sheet], new[sheet]
    oh, nh = header_index(o), header_index(n)
    saved = {}
    for row in o.iter_rows(min_row=2, values_only=True):
        call = row[oh["Your call"]]
        note = row[oh["Notes"]]
        if note and str(note).startswith(GENERATED_NOTE_PREFIXES):
            note = None
        if call or note:
            saved[(row[oh[key_col]], row[oh[desc_col]])] = (call, note)

    applied = 0
    for row in n.iter_rows(min_row=2):
        key = (row[nh[key_col]].value, row[nh[desc_col]].value)
        if key not in saved:
            continue
        call, note = saved.pop(key)
        if call:
            row[nh["Your call"]].value = call
        generated = row[nh["Notes"]].value
        if note and note != generated:
            row[nh["Notes"]].value = f"{generated} | {note}" if generated else note
        applied += 1
    print(f"{sheet}: carried over {applied}, not matched {len(saved)}")
    for (k, d), (call, note) in saved.items():
        print(f"   unmatched: {k} | {d} | {call} | {note}")

new.save(new_path)
